"""Export the dataset workbook letterboxd_reviews_dataset.xlsx.

Sheets:
  1. Reviews       - all sampled reviews (target 40 x 500 = 20,000 rows)
  2. Movies        - one row per film with population counts and status
  3. Sampling_Log  - per-film sampling statistics
  4. Methodology   - data source, procedure, ethics, limitations
  5. Validation    - results of the automated validation checks

Review texts are written as literal strings (never as formulas), and
characters that are illegal in XLSX are removed in the Excel copy only.
The CSV remains the authoritative, unmodified copy of the text.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parent))

import config

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
WRAP = Alignment(wrap_text=True, vertical="top")

WIDTHS = {
    "review_text": 80, "movie_title": 32, "review_url": 45, "reviewer_id": 25,
    "sampling_method": 40, "notes": 60, "source": 35, "Section": 28,
    "Description": 110, "check": 45, "detail": 70, "collected_at": 22,
}


def _clean_value(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(value) > config.EXCEL_MAX_CELL_CHARS:
            value = value[: config.EXCEL_MAX_CELL_CHARS - 20] + " [TRUNCATED IN XLSX]"
    return value


def _write_sheet(ws, df: pd.DataFrame, wrap_cols: tuple[str, ...] = ()) -> int:
    truncated = 0
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in df.itertuples(index=False):
        values = []
        for v in row:
            if isinstance(v, str) and len(v) > config.EXCEL_MAX_CELL_CHARS:
                truncated += 1
            values.append(_clean_value(v))
        ws.append(values)
        for cell in ws[ws.max_row]:
            # Force text: a review starting with "=" must not become a formula.
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for idx, col in enumerate(df.columns, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = WIDTHS.get(col, max(12, min(30, len(col) + 4)))
        if col in wrap_cols:
            for cell in ws[letter][1:]:
                cell.alignment = WRAP
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions
    return truncated


def export_excel(final: pd.DataFrame, movies_sheet: pd.DataFrame, sampling_log: pd.DataFrame,
                 methodology: list[tuple[str, str]], validation: pd.DataFrame,
                 path: Path | None = None) -> tuple[Path, int]:
    path = path or config.OUTPUT_DIR / config.XLSX_OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Reviews"
    truncated = _write_sheet(ws, final)
    _write_sheet(wb.create_sheet("Movies"), movies_sheet)
    _write_sheet(wb.create_sheet("Sampling_Log"), sampling_log, wrap_cols=("notes",))
    meth = pd.DataFrame(methodology, columns=["Section", "Description"])
    _write_sheet(wb.create_sheet("Methodology"), meth, wrap_cols=("Description",))
    _write_sheet(wb.create_sheet("Validation"), validation, wrap_cols=("detail",))
    wb.save(path)
    return path, truncated
