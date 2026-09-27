"""Tests of the pipeline on SYNTHETIC data only.

The synthetic reviews generated here exist solely to prove the pipeline
logic works; they are written to a temporary directory and never become
part of the thesis dataset.
"""
import random
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

import config  # noqa: E402
import pipeline  # noqa: E402
from ingest import hash_username, parse_rating  # noqa: E402

ADJ = ["brilliant", "boring", "beautiful", "confusing", "hilarious", "painful", "gorgeous", "clumsy",
       "moving", "forgettable", "ambitious", "tender", "loud", "clever", "messy", "haunting"]
NOUN = ["soundtrack", "ending", "script", "cinematography", "lead performance", "pacing", "dialogue",
        "villain", "editing", "final act", "production design", "score"]
VERB = ["loved", "hated", "admired", "enjoyed", "respected", "could not stand", "appreciated"]
ES = ["La película me pareció muy aburrida y demasiado larga para mi gusto, número {i}.",
      "Una obra maestra absoluta, la fotografía es preciosa y los actores están increíbles {i}."]
FR = ["Ce film est vraiment magnifique, j'ai adoré la musique et les acteurs, numéro {i}."]


def english(i: int, rng: random.Random) -> str:
    return (f"Honestly the {rng.choice(NOUN)} was {rng.choice(ADJ)} and the {rng.choice(NOUN)} felt "
            f"{rng.choice(ADJ)}. I {rng.choice(VERB)} watching this with my friends, review number {i}.")


def make_rows(movie_id, slug, n_en, n_other, rng, start=0):
    rows = []
    for i in range(start, start + n_en):
        rows.append({"movie_id": movie_id, "review_id": f"{movie_id}-{i:05d}",
                     "review_url": f"https://letterboxd.com/user{i}/film/{slug}/",
                     "review_text": english(i, rng), "review_date": f"2024-{1 + i % 12:02d}-{1 + i % 28:02d}",
                     "rating": rng.choice(["★★★", "★★★★½", "2.5", "5", ""]), "username": f"user{i}"})
    for j in range(n_other):
        tmpl = (ES + FR)[j % 3]
        rows.append({"movie_id": movie_id, "review_id": f"{movie_id}-X{j:04d}",
                     "review_url": f"https://letterboxd.com/intl{j}/film/{slug}/",
                     "review_text": tmpl.format(i=j), "review_date": "2024-06-01", "rating": "4",
                     "username": f"intl{j}"})
    return rows


@pytest.fixture
def sandbox(tmp_path, monkeypatch):
    for name in ("PROCESSED_DIR", "OUTPUT_DIR", "LOGS_DIR", "REPORTS_DIR"):
        monkeypatch.setattr(config, name, tmp_path / name.lower())
    raw = tmp_path / "raw"
    raw.mkdir()
    rng = random.Random(0)
    m01 = make_rows("M01", "the-godfather", 700, 30, rng)
    # Invalid and duplicate records for M01
    m01 += [{"movie_id": "M01", "review_id": "M01-E1", "review_url": "", "review_text": "🔥🔥🔥",
             "review_date": "2024-01-01", "rating": "", "username": "e1"},
            {"movie_id": "M01", "review_id": "M01-E2", "review_url": "", "review_text": "   ",
             "review_date": "2024-01-01", "rating": "", "username": "e2"},
            dict(m01[0]),  # duplicate review_id
            {**m01[1], "review_id": "M01-COPY", "review_url": ""},  # duplicate text
            {"movie_id": "M01", "review_id": "M01-F", "review_url": "",
             "review_text": "=HYPERLINK(\"x\") this formula-looking review should stay plain text in Excel.",
             "review_date": "2024-02-02", "rating": "3", "username": "formula"}]
    m02 = make_rows("M02", "pulp-fiction", 300, 10, rng, start=10000)  # insufficient
    pd.DataFrame(m01[:400]).to_csv(raw / "m01_part1.csv", index=False)
    pd.DataFrame(m01[400:]).to_csv(raw / "m01_part2.csv", index=False)
    df2 = pd.DataFrame(m02).drop(columns=["movie_id"])
    df2["letterboxd_slug"] = "pulp-fiction"
    df2.to_json(raw / "m02.jsonl", orient="records", lines=True, force_ascii=False)
    pd.DataFrame([{"movie_id": "UNKNOWN", "review_text": "orphan row", "username": "x"}]).to_csv(raw / "orphans.csv", index=False)
    pd.DataFrame([
        {"file": f, "source": "SYNTHETIC TEST DATA", "collected_at": "2026-09-27T00:00:00Z"}
        for f in ("m01_part1.csv", "m01_part2.csv", "m02.jsonl", "orphans.csv")
    ]).to_csv(raw / "manifest.csv", index=False)
    pd.DataFrame([{"movie_id": "M01", "total_reviews_available": "740"}]).to_csv(raw / "frame_sizes.csv", index=False)
    return tmp_path, raw


def test_full_pipeline(sandbox):
    tmp, raw = sandbox
    result = pipeline.run(raw_dir=raw, make_pdf=True)
    log = result["log"].set_index("movie_id")
    final = result["final"]

    assert log.loc["M01", "collection_status"] == config.STATUS_COMPLETE
    assert log.loc["M01", "sample_obtained"] == 500
    assert log.loc["M01", "invalid_removed"] == 2
    assert log.loc["M01", "non_english_reviews"] == 30
    assert log.loc["M01", "duplicates_removed"] == 2
    assert log.loc["M01", "eligible_population"] == 701
    assert log.loc["M01", "total_reviews_available"] == 740
    assert log.loc["M02", "collection_status"] == config.STATUS_INSUFFICIENT
    assert log.loc["M02", "sample_obtained"] == 300
    assert log.loc["M03", "collection_status"] == config.STATUS_AWAITING

    assert len(final) == 800
    assert set(final["language"]) == {"en"}
    assert not final["review_id"].duplicated().any()
    assert (final["sampling_seed"] == 42).all()
    assert list(final.columns) == config.FINAL_COLUMNS
    assert result["summary"]["integrity_checks_passed"]
    assert not result["summary"]["dataset_complete"]
    assert result["totals"]["rejected_rows"] == 1

    # No usernames anywhere in outputs
    for f in list((tmp / "output_dir").glob("*.csv")) + list((tmp / "processed_dir").glob("*.csv")):
        assert "username" not in pd.read_csv(f, nrows=0).columns
    assert (tmp / "reports_dir" / config.PDF_REPORT_NAME).stat().st_size > 5000

    wb = load_workbook(tmp / "output_dir" / config.XLSX_OUTPUT_NAME)
    assert wb.sheetnames[:4] == ["Reviews", "Movies", "Sampling_Log", "Methodology"]
    assert wb["Reviews"].max_row == 801
    assert wb["Movies"].max_row == 41


def test_sampling_is_reproducible_and_order_independent(sandbox):
    tmp, raw = sandbox
    first = set(pipeline.run(raw_dir=raw, make_pdf=False)["final"]["review_id"])
    # Rename files so they are read in a different order; sample must not change.
    (raw / "m01_part1.csv").rename(raw / "z_m01_part1.csv")
    manifest = pd.read_csv(raw / "manifest.csv")
    manifest["file"] = manifest["file"].replace({"m01_part1.csv": "z_m01_part1.csv"})
    manifest.to_csv(raw / "manifest.csv", index=False)
    second = set(pipeline.run(raw_dir=raw, make_pdf=False)["final"]["review_id"])
    assert first == second


def test_formula_text_stays_text_in_excel(sandbox):
    tmp, raw = sandbox
    pipeline.run(raw_dir=raw, make_pdf=False)
    ws = load_workbook(tmp / "output_dir" / config.XLSX_OUTPUT_NAME)["Reviews"]
    for row in ws.iter_rows(min_row=2):
        assert row[4].data_type == "s"


def test_rating_and_hash_helpers():
    assert parse_rating("★★★★½") == 4.5
    assert parse_rating("8/10") == 4.0
    assert parse_rating("3.5") == 3.5
    assert parse_rating("") is None
    assert parse_rating("7") is None
    assert hash_username("Alice", salt="") == hash_username(" alice ", salt="")
    assert hash_username("alice", salt="s1") != hash_username("alice", salt="s2")
    assert hash_username("", salt="") == ""
